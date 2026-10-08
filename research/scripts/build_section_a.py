"""Build Section A sheets (market size, firm concentration, growth drivers) into the workbook."""
import json, sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

path = sys.argv[1]
wb = load_workbook(path)
for n in ('A. Market size', 'A. Firms', 'A. Growth drivers'):
    if n in wb.sheetnames: del wb[n]

H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864')
T = Font(bold=True, size=14); S = Font(bold=True, size=11); N = Font(italic=True, size=9)
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

def table(ws, r, headers, rows, widths=None):
    for c, h in enumerate(headers, 1):
        x = ws.cell(r, c, h); x.font = H; x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical='center'); x.border = BOX
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            x = ws.cell(r, c, v); x.border = BOX; x.alignment = WRAP
    if widths:
        for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
    return r + 2

# ---------------- A. Market size ----------------
ws = wb.create_sheet('A. Market size', 0)
ws['A1'] = 'Section A. Swedish (corporate) legal market: size'; ws['A1'].font = T
ws['A3'] = 'A1. Headline estimates'; ws['A3'].font = S
r = table(ws, 4, ['Scope', 'Figure', 'Year', 'Source'], [
    ['Business law: fee income, 56 largest firms (Affärsvärlden panel)', '~SEK 16bn (+5.3%)', '2025', 'https://www.affarsvarlden.se/artikel/affarsjuridiken-vaxer-igen-men-jattarna-tappar-mark'],
    ['Same panel, prior year (revised Oct-25, split FYs)', 'SEK 15.4bn (+8.6%); SEK 4.13m fee per lawyer', '2024', 'https://www.realtid.se/affarer/154-miljarder-sa-klarade-affarsjuristerna-tullchocken/'],
    ['Business law market, total (industry estimate, method not given)', '>SEK 25bn p.a.', 'n/a', 'https://www.affarsvarlden.se/debatt/okad-konkurrens-starker-svensk-affarsjuridik'],
    ['All legal services (MarketLine)', 'USD 3.1bn (CAGR 5.9% 2019-24)', '2024', 'https://store.marketline.com/report/legal-services-in-sweden/'],
    ['All legal activities (IBISWorld)', 'EUR 3.5bn (CAGR -0.5% 2020-25)', '2026e', 'https://www.ibisworld.com/sweden/industry/legal-activities/200283/'],
    ['All law firms (advokatbyråer), estimate from SCB', '~SEK 29bn', '2024e', '74.5% share of SNI 69.1 in 2021 applied to 2024 (SCB, below)'],
], [58, 40, 9, 70])

ws.cell(r, 1, 'A2. Official statistics: SCB Företagens ekonomi, net turnover (SEK m)').font = S
rows = [[2010, 18738, 13904, 1956, 2879], [2015, 23986, 17936, 2750, 3300], [2019, 31237, 22277, 4194, 4765],
        [2020, 32628, 23298, 4371, 4958], [2021, 34693, 25864, 4558, 4270],
        ['2022*', 33276, 'n/p', 'n/p', 'n/p'], ['2023*', 36525, 'n/p', 'n/p', 'n/p'], ['2024*', 38921, 'n/p', 'n/p', 'n/p']]
r = table(ws, r + 1, ['Year', '69.1 Legal activities (total)', '69.101 Advokatbyråer', '69.102 Other legal offices', '69.103 Patent agents'], rows)
for rr in range(r - len(rows) - 2, r - 1):
    for c in range(2, 6): ws.cell(rr, c).number_format = '#,##0'
for t in ['* New SCB series from 2022 (NV0109P): 2022 onward is not directly comparable to 2021; sub-industry splits not published (n/p).',
          '69.1 in 2024: 6,965 companies, 13,446 employees; production value SEK 39.1bn, value added SEK 25.3bn, EBITDA margin 28.2%.',
          '69.101 turnover CAGR 2010-21: ~5.8%. Source: https://api.scb.se/OV0104/v1/doris/sv/ssd/NV/NV0109/ (BasfaktaFEngs07, NSEBasfaktaFEngs07, NSEBasfaktaVEngs07N).']:
    ws.cell(r - 1, 1, t).font = N; r += 1
ws.column_dimensions['B'].width = 40; ws.column_dimensions['C'].width = 20; ws.column_dimensions['D'].width = 24; ws.column_dimensions['E'].width = 20

# ---------------- A. Firms ----------------
# firm, FY, revenue SEKm, revenue basis, partners, partner src, lawyers, employees, employee src, notes
F = [
    ['Vinge', '2025', 2036, 'Group, annual reports (DJ via Realtid)', 63, 'Dividend recipients: Stockholm 49 + Malmö 14; excl. Göteborg (understated)', None, 550, 'Website, 2026 (~550)', '-2% y/y; 2024 was a record (2,083-2,085)'],
    ['Mannheimer Swartling', '2025', 2002, 'Group, annual report (DJ via Realtid)', 92, 'DJ (dividend split), 2025', 430, 650, 'Chambers 2026 / AB 649', '+4%; EBIT SEK 796m; dividend SEK 617m'],
    ['Roschier (Sweden)', 'FY Jun24-May25', 1254, 'Affärsvärlden Oct-25', None, 'Group 52 incl. Finland (not used)', None, 226, 'Company accounts (AB)', 'FY25/26 net sales 1,304 (220 staff)'],
    ['White & Case (Stockholm)', '2025', 1108, 'Company accounts, net sales (AB)', 16, 'Implied (est.): DJ ">SEK 69m per partner"', 70, 120, 'Company accounts (AB)', '2024: 859; +29%. AB net sales may include intra-group billing; per-head ratios look high, treat with caution'],
    ['Setterwalls', '2025', 1000, 'Affärsvärlden (">1bn"); parent AB 1,056', 37, 'Chambers 2026 (firm-supplied; looks low)', None, 306, 'Sum of 3 entities (AB)', 'Malmö office (SEK 305m) left Sep-26 for AGRD'],
    ['Lindahl', '2024', 737, 'Affärsvärlden Apr-25', None, '', 153, 300, 'Chambers 2026', '2025 entity sum ~635 (est., not comparable)'],
    ['Gernandt & Danielsson', 'FY24/25', 676, 'Affärsvärlden Oct-25', 21, 'Chambers 2026', 80, 150, 'Firm website', '+12% (FY23/24: 601)'],
    ['DLA Piper (Sweden)', 'FY24/25', 602, 'Affärsvärlden Oct-25', None, '', None, None, '', '-9%'],
    ['CMS Wistrand', '2024', 558, 'Affärsvärlden Apr-25', None, '', None, None, '', ''],
    ['Schjødt (Sweden)', '2025', 534, 'Est.: ">SEK 0.5bn", +25% on 427', None, '', None, None, '', 'Branch, no Swedish accounts; 8 Sthlm partners left for BAHR (2026)'],
    ['Delphi', '2024', 516, 'Affärsvärlden Apr-25', 49, 'Chambers 2026', 150, 220, 'Firm website', '+2%'],
    ['Cirio', '2025', 431, 'Company accounts, net sales (AB)', None, '', None, 127, 'Company accounts (AB)', '+15% (Afv); 2024 Afv 363'],
    ['MAQS', '2024', 417, 'Affärsvärlden Apr-25', None, '', 127, 180, 'Chambers 2026', '+21%'],
    ['Baker McKenzie (Stockholm)', 'FY23/24', 334, 'Affärsvärlden Apr-25', None, '', None, None, '', 'KB, no public accounts'],
    ['Snellman (Sweden)', '2025', 279, 'Company accounts, net sales (AB)', None, '', None, 107, 'Company accounts (AB)', '2024 Afv 295'],
    ['Foyen', '2024', 277, 'Affärsvärlden Apr-25', None, '', None, 106, 'Company accounts (AB)', ''],
    ['Glimstedt', '2024', 196, 'Affärsvärlden Apr-25', None, '', None, None, '', '~200 staff on website includes Baltics (not used)'],
    ['Eversheds Sutherland', '2025', 166, 'Company accounts, net sales (AB)', None, '', None, 59, 'Company accounts (AB)', '2024 Afv 140'],
    ['Fylgia', '2024', 131, 'Affärsvärlden Apr-25', None, '', None, None, '', 'Record 2025 (firm)'],
    ['Westerberg & Partners', '2025', 107, 'Company accounts, net sales (AB)', 12, 'Chambers 2026', None, 35, 'Company accounts (AB)', ''],
    ['Moll Wendén', '2025', 101, 'Company accounts, net sales (AB)', None, '', None, 44, 'Company accounts (AB)', ''],
    ['Born', '2024', 86, 'Affärsvärlden Apr-25', None, '', None, None, '', 'Part of AGRD since 2025'],
    ['Harvest', '2025', 82, 'Company accounts, net sales (AB)', None, '', None, 27, 'Company accounts (AB)', ''],
    ['Synch', '2025', 81, 'Company accounts, net sales (AB)', None, '', None, 38, 'Company accounts (AB)', 'Part of AGRD since 2025'],
    ['Kastell', '2025', 80, 'Company accounts, net sales (AB)', 9, 'Chambers 2026', 19, 21, 'Company accounts (AB)', ''],
]
F.sort(key=lambda x: -x[2])
ws = wb.create_sheet('A. Firms', 1)
ws['A1'] = 'A3. Concentration: largest Swedish business law firms ranked by latest revenue'; ws['A1'].font = T
ws['A2'] = ('Revenue = latest available year per firm (SEK m). Affärsvärlden figures are fee income; company accounts (AB) are net sales of the main legal entity and are not '
            'strictly comparable. Partners and lawyers mostly firm-supplied to Chambers (2026). Employees = headcount, used as a proxy for FTEs (no FTE data is published); '
            'AB headcounts usually exclude partners who bill via own companies. Ratios are calculated only where both inputs exist. Grey = not found.')
ws['A2'].alignment = WRAP; ws.merge_cells('A2:N2'); ws.row_dimensions[2].height = 55
hdr = ['#', 'Firm', 'FY', 'Revenue SEK m', 'Revenue basis', 'Partners', 'Partner source', 'Lawyers', 'Employees (FTE proxy)', 'Employee source',
       'Rev / partner SEK m', 'Rev / employee SEK m', 'Rev / lawyer SEK m', 'Notes']
R = 4
for c, h in enumerate(hdr, 1):
    x = ws.cell(R, c, h); x.font = H; x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); x.border = BOX
grey = PatternFill('solid', fgColor='EDEDED')
for i, f in enumerate(F, 1):
    r = R + i
    vals = [i, f[0], f[1], f[2], f[3], f[4], f[5], f[6], f[7], f[8]]
    for c, v in enumerate(vals, 1):
        x = ws.cell(r, c, v); x.border = BOX; x.alignment = Alignment(vertical='top', wrap_text=c in (5, 7, 10))
        if v in (None, '') and c in (6, 8, 9): x.fill = grey
    ws.cell(r, 4).number_format = '#,##0'
    for c, num in ((11, 'F'), (12, 'I'), (13, 'H')):
        x = ws.cell(r, c); x.border = BOX
        if ws[f'{num}{r}'].value:
            x.value = f'=D{r}/{num}{r}'; x.number_format = '0.0'
        else:
            x.fill = grey
    x = ws.cell(r, 14, f[9]); x.border = BOX; x.alignment = WRAP
last = R + len(F)
r = last + 1
ws.cell(r, 2, 'Total shown').font = Font(bold=True); ws.cell(r, 4, f'=SUM(D{R+1}:D{last})').number_format = '#,##0'
ws.cell(r + 1, 2, 'Top 5 share of ~SEK 16bn panel (indicative, mixed years/bases)'); ws.cell(r + 1, 4, f'=SUM(D{R+1}:D{R+5})/16000').number_format = '0%'
ws.cell(r + 2, 2, 'Top 10 share of ~SEK 16bn panel (indicative)'); ws.cell(r + 2, 4, f'=SUM(D{R+1}:D{R+10})/16000').number_format = '0%'
notes = [
    'Caveats: Vinge partner count covers only the Stockholm and Malmö dividend-receiving partners (no Göteborg), so Vinge rev/partner is overstated. White & Case partners are implied from Dagens Juridik ">SEK 69m per partner" (paywalled).',
    'Setterwalls partner count (37) is as supplied to Chambers and looks low against the firm size. Roschier group partner count (52) includes Finland and is not used. Magnusson: no Swedish revenue or headcount found.',
    'Sources: Affärsvärlden 12-Apr-25 https://www.affarsvarlden.se/artikel/vinge-gar-om-msa-sa-gar-det-for-affarsjuristerna ; 3-Oct-25 https://www.affarsvarlden.se/artikel/affarsjuristerna-gar-starkt-finns-fler-bolag-i-ipo-pipelinen ;',
    'Realtid 7-Aug-26 https://www.realtid.se/juridik/vinge-tappar-mark-mannheimer-swartling-drar-ifran/ ; Realtid 26-Apr-26 https://www.realtid.se/juridik/affarsjuridiken-vaxer-men-sprickorna-syns/ ; Realtid 27-Nov-25 https://www.realtid.se/juridik/white-case-dubblade-omsattningen-var-styrka/ ;',
    'Realtid 1-Sep-26 https://www.realtid.se/juridik/setterwalls-malmokontor-gar-till-riskkapitalagda-agrd/ ; Affärsvärlden 5-Oct-26 (Schjødt) https://www.affarsvarlden.se/artikel/schjdt-vaxer-snabbast--men-konkurrensen-hardnar ;',
    'company accounts via https://www.allabolag.se/{org.nr} ; Chambers Europe 2026 firm profiles (firm-supplied "top figures"); firm websites.',
]
for k, t in enumerate(notes):
    ws.cell(r + 4 + k, 1, t).font = N
for c, w in zip(range(1, 15), (4, 26, 13, 11, 30, 9, 30, 9, 11, 24, 10, 10, 10, 40)):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = 'C5'

# ---------------- A. Growth drivers ----------------
ws = wb.create_sheet('A. Growth drivers', 2)
ws['A1'] = 'A4. Fastest growers and growth drivers'; ws['A1'].font = T
ws['A2'] = 'Driver tags: Lateral hire / M&A / New office / Organic. Only drivers found in a cited source are listed.'; ws['A2'].font = N
G = [
    ['Carthiel', '+155%', '2024', 'M&A', 'Merged with Grey Advokatbyrå (operating from Feb-2024), adding a disputes team. Small base (~10 people).', 'https://www.carthiel.se/post/advokatfirman-carthiel-g%C3%A5r-ihop-med-grey-advokatbyr%C3%A5'],
    ['White & Case (Stockholm)', '+29% (AB net sales 859 to 1,108)', '2025', 'Organic', 'Fastest grower in 2023 (+24%) on PE deals and selective hiring; "doubled revenue" coverage Nov-25. Lost Magnus Wennerhorn to Linklaters (Sep-24).', 'https://www.realtid.se/juridik/white-case-dubblade-omsattningen-var-styrka/'],
    ['Schjødt (Sweden)', '+24% / ~+25%', '2024 / 2025', 'Lateral hire', 'Tom Wehtje (M&A) joined 2024 from Mannheimer Swartling after 15 years as partner; Karl Klackenberg (M&A/PE) reportedly from Vinge May-24 (unconfirmed).', 'https://www.legal500.com/firms/10502-advokatfirmaet-schjodt-as/c-sweden/lawyers/2101480-tom-wehtje'],
    ['', '', '', 'Organic', 'Demand for pan-Scandinavian cross-border offering, FDI and regulatory work.', 'https://www.affarsvarlden.se/artikel/vinge-gar-om-msa-sa-gar-det-for-affarsjuristerna'],
    ['', '', '', 'M&A (historic)', 'Entered Sweden via merger with Hamilton (2019/20). Risk: 54 people (23 partners, group-wide) left for BAHR, which opened Stockholm Jan-26.', 'https://www.realtid.se/juridik/oppnar-stockholmskontor-efter-massvarvning-fran-rivalen/'],
    ['Roschier (Sweden)', '+23% / +22%', 'FY23/24 / FY24/25', 'Organic', 'Swedish transaction market pick-up, larger mandates, high disputes demand, closer Finland-Sweden integration; Northvolt Expansion Ett restructuring. Stockholm now larger than Helsinki.', 'https://www.realtid.se/affarer/154-miljarder-sa-klarade-affarsjuristerna-tullchocken/'],
    ['MAQS', '+21%', '2024', 'Organic', 'Strong demand in M&A, capital markets and disputes (CEO Katrin Troedsson).', 'https://www.affarsvarlden.se/artikel/vinge-gar-om-msa-sa-gar-det-for-affarsjuristerna'],
    ['', '', '', 'Lateral hire', 'Strategic sustainability hires, e.g. Anna Domander (ex head of legal, Systembolaget); date not given.', 'https://www.affarsvarlden.se/artikel/vinge-gar-om-msa-sa-gar-det-for-affarsjuristerna'],
    ['Vinge', '+19%', '2024', 'Organic', 'Broad client activity and full-service model; accounted for ~25% of total market growth. Then -2% in 2025 ("2024 exceptional").', 'https://www.affarsvarlden.se/artikel/vinge-gar-om-msa-sa-gar-det-for-affarsjuristerna'],
    ['Gernandt & Danielsson', '+19% / +12%', 'FY23/24 / FY24/25', 'Lateral hire', 'Linklaters Stockholm real estate team: partners Magnus Lidman and Anna Eriksson + 5 associates (dated Jun-24 by Global Legal Post/Legal 500; Afv says early 2025).', 'https://www.affarsvarlden.se/artikel/affarsjuristerna-gar-starkt-finns-fler-bolag-i-ipo-pipelinen'],
    ['', '', '', 'Organic', 'Strong bond and M&A teams.', 'https://www.affarsvarlden.se/artikel/affarsjuristerna-gar-starkt-finns-fler-bolag-i-ipo-pipelinen'],
    ['Cirio', '~+15%', '2025', 'Lateral hire', '12-person real estate and construction team from AG Advokat from Jan-25, incl. partners Johan Lindberg, Johan Rosén, Måns Derk; new construction disputes offering.', 'https://www.realtid.se/juridik/cirio-expanderar-inom-fastighets-och-entreprenadjuridik/'],
    ['', '', '', 'Organic', 'Investment in real estate transactions and insolvency.', 'https://www.realtid.se/juridik/affarsjuridiken-vaxer-men-sprickorna-syns/'],
    ['Setterwalls', '+2% / ~+9% (passed SEK 1bn)', '2024 / 2025', 'Organic', 'Large complex mandates: public M&A and PE take-privates (MedHelp, Integrum, Norva24), bank finance, disputes; lawyers +10% to 235 in 2024. Malmö office left for AGRD in Sep-26.', 'https://www.realtid.se/affarer/154-miljarder-sa-klarade-affarsjuristerna-tullchocken/'],
    ['Fylgia', 'n/a (record 2025)', '2025', 'Organic', 'Transactions, construction/real estate, public procurement, restructuring/insolvency; four new partners appointed.', 'https://www.fylgia.se/en/news/following-a-record-breaking-year-fylgia-appoints-four-new-partners'],
    ['AGRD Partners (Axcel)', 'n/a (~SEK 660m combined)', '2025', 'M&A', 'PE-backed roll-up of six firms in Aug-25: Allié, Born, Morris Law, Next, Synch, TM & Partners (~130 lawyers gave up advokat title). Added Setterwalls Malmö (Ascente Law) Sep-26.', 'https://www.realtid.se/juridik/setterwalls-malmokontor-gar-till-riskkapitalagda-agrd/'],
    ['DLA Piper (Sweden)', '-9%', 'FY24/25', 'Decline', 'Consolidation year: generational shift and hesitant deal market; restructuring group busy (Northvolt, Oscar Properties, SAS).', 'https://www.affarsvarlden.se/artikel/affarsjuristerna-gar-starkt-finns-fler-bolag-i-ipo-pipelinen'],
]
table(ws, 4, ['Firm', 'Growth', 'Period', 'Driver', 'Detail', 'Source'], G, [24, 22, 16, 14, 80, 60])
ws.cell(6 + len(G), 1, 'Market context: 2025 growth came from international transactions, disputes and regulatory work; IPOs were the weakest segment. Firms expect +6.8% in 2026 (Affärsvärlden/Realtid).').font = N

wb.move_sheet('A. Market size', -wb.index(wb['A. Market size']))
order = ['A. Market size', 'A. Firms', 'A. Growth drivers', 'B. Practice rankings', 'B. Rankings raw']
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save(path)
print('ok', wb.sheetnames)
