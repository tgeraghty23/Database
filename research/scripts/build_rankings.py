"""Build Section B (Chambers & Legal 500 Sweden practice rankings) workbook sheets."""
import json, re, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = 'research/data/'
ch = json.load(open(D + 'chambers_europe_2026_sweden_firm_bands.json'))
l5 = json.load(open(D + 'legal500_emea_sweden_firm_tiers.json'))

CANON = {
    'Advokatfirman Vinge KB': 'Vinge', 'Mannheimer Swartling': 'Mannheimer Swartling', 'Roschier': 'Roschier',
    'Gernandt & Danielsson Advokatbyrå KB': 'Gernandt & Danielsson', 'Gernandt & Danielsson Advokatbyrå': 'Gernandt & Danielsson',
    'Advokatfirman Cederquist KB': 'Cederquist', 'Setterwalls': 'Setterwalls', 'Snellman Advokatbyrå AB': 'Snellman', 'Snellman': 'Snellman',
    'Linklaters Advokatbyrå': 'Linklaters', 'Linklaters': 'Linklaters', 'White & Case Advokat AB': 'White & Case',
    'Baker McKenzie': 'Baker McKenzie', 'DLA Piper': 'DLA Piper', 'DLA Piper Sweden': 'DLA Piper',
    'Advokatfirman Hammarskiöld & Co AB': 'Hammarskiöld', 'Hammarskiöld': 'Hammarskiöld',
    'Advokatfirman Lindahl': 'Lindahl', 'Advokatfirman Lindahl KB': 'Lindahl', 'Delphi': 'Delphi', 'Advokatfirman Delphi': 'Delphi',
    'Magnusson': 'Magnusson', 'MAQS Advokatbyrå': 'MAQS', 'Advokatfirman Schjødt': 'Schjødt', 'Cirio Advokatbyrå AB': 'Cirio',
    'CMS Wistrand': 'CMS Wistrand', 'BAHR': 'BAHR', 'Bird & Bird': 'Bird & Bird', 'Eversheds Sutherland Advokatbyrå AB': 'Eversheds Sutherland',
    'Eversheds Sutherland': 'Eversheds Sutherland', 'KANTER Advokatbyrå': 'Kanter', 'TM & Partners': 'TM & Partners',
    'Wigge': 'Wigge & Partners', 'Wigge & Partners': 'Wigge & Partners', 'Advokatfirman Fylgia KB': 'Fylgia',
    'Advokatfirman Kahn Pedersen': 'Kahn Pedersen', 'Advokatfirman Kahn Pedersen KB': 'Kahn Pedersen',
    'BOKWALL RISLUND Advisors': 'Bokwall Rislund', 'BOKWALL RISLUND': 'Bokwall Rislund',
    'Sandart & Partners': 'Sandart & Partners', 'Sandart&Partners Advokatbyrå KB': 'Sandart & Partners',
    'Westerberg & Partners': 'Westerberg & Partners', 'Kastell Advokatbyrå': 'Kastell', 'Kastell Advokatbyrå AB': 'Kastell',
    'Morris Law AB': 'Morris Law', 'Morris Law': 'Morris Law', 'Next Law KB': 'Next Law', 'RE:FI STHLM': 'RE:FI STHLM',
    'Fröberg & Lundholm Advokatbyrå': 'Fröberg & Lundholm', 'Norburg & Scherp Advokatbyrå': 'Norburg & Scherp', 'Norburg & Scherp': 'Norburg & Scherp',
    'Advokatbyrån Gulliksson AB': 'Gulliksson', 'Advokatbyrån Gulliksson': 'Gulliksson', 'Wallin & Partners': 'Wallin & Partners',
    'Advokatbyrån Wallin & Partners AB': 'Wallin & Partners', 'Brick Advokat': 'Brick', 'Brick Advokat AB': 'Brick',
    'Elmzell Advokatbyrå': 'Elmzell', 'Elmzell Advokatbyrå AB, member of Ius Laboris': 'Elmzell', 'Kompass Advokat': 'Kompass',
    'REAL Advokatbyrå': 'REAL Advokatbyrå', 'Real Advokatbyrå Aktiebolag': 'REAL Advokatbyrå', 'Ström Advokatbyrå': 'Ström',
    'Skierfe Advokatfirma': 'Skierfe', 'Skierfe Advokatfirma KB': 'Skierfe', 'Advokatfirman Ulfsdotter AB': 'Ulfsdotter',
    'Advokatfirman Ulfsdotter AB/Ulfsdotter Law': 'Ulfsdotter', 'Hellström Advokatbyrå kb': 'Hellström', 'Hellström Law': 'Hellström',
    'Synch Law AB': 'Synch', 'Born': 'Born', 'Allié': 'Allié', 'Harvest Advokatbyrå': 'Harvest', 'Foyen Advokatfirma': 'Foyen',
    'AG Advokat': 'AG Advokat', 'Advokatfirman Carler': 'Carler', 'Edel Advokater': 'Edel', 'A1 Advokater': 'A1 Advokater',
    'Andulf Advokat AB': 'Andulf', 'NORMA Advokater KB': 'Norma', 'Merino Advokatbyrå': 'Merino', 'NORDIA Law': 'Nordia Law',
}
def canon(n): return CANON.get(n, n)

# (label, chambers key, legal500 key)
AREAS = [
    ('Corporate / M&A', 'corporate-ma', 'commercial-corporate-and-ma'),
    ('Private equity', 'private-equity', None),
    ('PE fund formation', 'private-equity-fund-formation', None),
    ('Banking & finance', 'banking-finance', 'banking-and-finance'),
    ('Capital markets: equity', 'capital-markets-equity', 'capital-markets-equity'),
    ('Capital markets: debt', 'capital-markets-debt', 'capital-markets-debt'),
    ('Dispute resolution', 'dispute-resolution', 'dispute-resolution'),
    ('Restructuring / insolvency', 'restructuring-insolvency', 'insolvency'),
    ('Insolvency trustees', 'restructuring-insolvency-trustees-administrators', None),
    ('Competition / EU', 'competition-european-law', 'eu-and-competition'),
    ('Employment', 'employment', 'employment'),
    ('Tax', 'tax', 'tax'),
    ('Real estate', 'real-estate', 'real-estate'),
    ('Construction', None, 'construction'),
    ('IP', 'intellectual-property', 'intellectual-property-and-media'),
    ('IT / TMT', 'information-technology', 'it-and-telecoms'),
    ('Data protection', None, 'data-privacy-and-data-protection'),
    ('Fintech', None, 'fintech'),
    ('Energy', 'energy-natural-resources', 'energy'),
    ('Environment', 'environment', 'environment'),
    ('Insurance', 'insurance', 'insurance'),
    ('Public procurement', 'public-procurement', 'public-procurement'),
    ('Healthcare / life sciences', None, 'healthcare-and-life-sciences'),
    ('Shipping', None, 'shipping'),
    ('Corporate investigations', 'corporate-investigations', None),
]

def short(label):
    if label.startswith('Band'): return 'B' + label[-1]
    if label.startswith('Tier'): return 'T' + label[-1]
    return 'FTW'  # Legal 500 'Firms to watch'

grid = {}  # firm -> {(area, dir): code}
raw = []
for area, ck, lk in AREAS:
    for dname, src, key in (('Chambers', ch, ck), ('Legal 500', l5, lk)):
        if not key: continue
        for label, n in src[key]:
            f = canon(n)
            grid.setdefault(f, {})[(area, dname)] = short(label)
            raw.append((area, dname, label, f, n))

def score(f):
    # points: band/tier 1 = 5 ... 5 = 1; firms to watch = 0.5
    s = 0
    for v in grid[f].values():
        s += 0.5 if v == 'FTW' else 6 - int(v[1])
    return s

BOX_SINGLE = ['MAQS']
BOX_GROUP = ['Magnusson', 'Delphi', 'Lindahl']
firms = sorted(grid, key=lambda f: -score(f))
firms = [f for f in firms if score(f) >= 4 or f in BOX_SINGLE + BOX_GROUP]
# keep the boxed group contiguous, placed where its highest-scoring member falls
pos = min(firms.index(f) for f in BOX_GROUP)
rest = [f for f in firms if f not in BOX_GROUP]
grp = sorted(BOX_GROUP, key=lambda f: -score(f))
firms = rest[:pos] + grp + rest[pos:]

path = sys.argv[1]
try:
    wb = load_workbook(path)
except FileNotFoundError:
    wb = Workbook(); wb.remove(wb.active)
for n in ('B. Practice rankings', 'B. Rankings raw'):
    if n in wb.sheetnames: del wb[n]
ws = wb.create_sheet('B. Practice rankings')

hdr = Font(bold=True, color='FFFFFF'); hfill = PatternFill('solid', fgColor='1F3864')
sub = PatternFill('solid', fgColor='D9E1F2')
fills = {'1': 'C6E0B4', '2': 'E2EFDA', '3': 'FFF2CC', '4': 'FCE4D6', '5': 'F2F2F2'}
thin = Side(style='thin', color='BFBFBF'); thick = Side(style='thick', color='C00000')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

ws['A1'] = 'Section B. Chambers & Legal 500 practice rankings: Sweden'
ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = ('Chambers Europe 2026 (firm bands, B1 = top) and Legal 500 EMEA, current edition (firm tiers, T1 = top; FTW = firm to watch). '
            'Extracted from chambers.com and legal500.com on 2026-10-08. Blank = not ranked. Score = sum of points (band/tier 1 = 5 ... 5 = 1, FTW = 0.5). '
            'Rows sorted by score; firms with score < 4 omitted, except boxed firms. Magnusson, Delphi and Lindahl are grouped together so they can share one box.')
ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4 + 2 * len(AREAS))
ws.row_dimensions[2].height = 45

R0 = 4
ws.cell(R0, 1, 'Firm'); ws.cell(R0, 2, 'Score'); ws.cell(R0, 3, '# Chambers'); ws.cell(R0, 4, '# Legal 500')
for c in range(1, 5):
    ws.merge_cells(start_row=R0, start_column=c, end_row=R0 + 1, end_column=c)
col = 5
for area, ck, lk in AREAS:
    ws.cell(R0, col, area); ws.merge_cells(start_row=R0, start_column=col, end_row=R0, end_column=col + 1)
    ws.cell(R0 + 1, col, 'CH' if ck else 'CH n/a'); ws.cell(R0 + 1, col + 1, 'L500' if lk else 'L500 n/a')
    col += 2
last = col - 1
for r in (R0, R0 + 1):
    for c in range(1, last + 1):
        cell = ws.cell(r, c); cell.font = hdr if r == R0 else Font(bold=True, size=8); cell.alignment = center
        cell.fill = hfill if r == R0 else sub
ws.row_dimensions[R0].height = 48

r = R0 + 2
rowof = {}
for f in firms:
    rowof[f] = r
    ws.cell(r, 1, f).font = Font(bold=f in BOX_SINGLE + BOX_GROUP)
    ws.cell(r, 2, score(f))
    ws.cell(r, 3, sum(1 for (a, d) in grid[f] if d == 'Chambers'))
    ws.cell(r, 4, sum(1 for (a, d) in grid[f] if d == 'Legal 500'))
    c = 5
    for area, ck, lk in AREAS:
        for i, d in enumerate(('Chambers', 'Legal 500')):
            v = grid[f].get((area, d))
            cell = ws.cell(r, c + i, v)
            cell.alignment = center
            if v and v != 'FTW': cell.fill = PatternFill('solid', fgColor=fills[v[1]])
            elif ((ck is None and i == 0) or (lk is None and i == 1)): cell.fill = PatternFill('solid', fgColor='EDEDED')
        c += 2
    for cc in range(1, last + 1):
        ws.cell(r, cc).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    r += 1

def box(r1, r2):
    for rr in range(r1, r2 + 1):
        for cc in range(1, last + 1):
            b = ws.cell(rr, cc).border
            ws.cell(rr, cc).border = Border(
                left=thick if cc == 1 else b.left, right=thick if cc == last else b.right,
                top=thick if rr == r1 else b.top, bottom=thick if rr == r2 else b.bottom)
box(rowof['MAQS'], rowof['MAQS'])
box(min(rowof[f] for f in BOX_GROUP), max(rowof[f] for f in BOX_GROUP))

ws.column_dimensions['A'].width = 24
for c in range(2, 5): ws.column_dimensions[get_column_letter(c)].width = 8
for c in range(5, last + 1): ws.column_dimensions[get_column_letter(c)].width = 6.5
ws.freeze_panes = ws.cell(R0 + 2, 5)
r += 1
ws.cell(r, 1, 'Not shown: Chambers tables that rank individual lawyers only (Banking & Finance: Regulatory, Data Protection, Telecommunications, '
              'Most in Demand Arbitrators). Legal 500 practice areas without a Chambers equivalent show "CH n/a" (and vice versa).').font = Font(italic=True, size=9)
r += 1
ws.cell(r, 1, 'Sources: https://chambers.com/legal-guide/europe-7 (Sweden) ; https://www.legal500.com/c/sweden/practice-areas').font = Font(italic=True, size=9)

wr = wb.create_sheet('B. Rankings raw')
wr.append(['Practice area', 'Directory', 'Band / tier', 'Firm (normalised)', 'Firm (as listed)'])
for c in wr[1]: c.font = Font(bold=True)
for row in raw: wr.append(row)
for c, w in zip('ABCDE', (26, 11, 14, 24, 40)): wr.column_dimensions[c].width = w
wr.auto_filter.ref = f'A1:E{len(raw) + 1}'
wb.save(path)
print('firms', len(firms), 'raw rows', len(raw)); print([(f, score(f)) for f in firms])
