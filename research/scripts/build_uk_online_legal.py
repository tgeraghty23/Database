"""Build the UK online / subscription legal services market map workbook from research/data/uk_online_legal_*.json."""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

out = sys.argv[1]
D = os.path.join(os.path.dirname(__file__), '..', 'data')
def load(n):
    with open(os.path.join(D, n), encoding='utf-8') as f: return json.load(f)

H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864')
T = Font(bold=True, size=14); S = Font(bold=True, size=11); N = Font(italic=True, size=9)
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
GREY = PatternFill('solid', fgColor='F2F2F2')

def j(v):
    if v is None: return None
    if isinstance(v, list): return '\n'.join(str(x) for x in v if x is not None)
    return v

def table(ws, r, headers, rows, widths=None, freeze=False):
    for c, h in enumerate(headers, 1):
        x = ws.cell(r, c, h); x.font = H; x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical='center'); x.border = BOX
    if freeze: ws.freeze_panes = ws.cell(r + 1, 3)
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            x = ws.cell(r, c, j(v)); x.border = BOX; x.alignment = WRAP
            if isinstance(v, (int, float)) and not isinstance(v, bool): x.number_format = '#,##0.0' if isinstance(v, float) else '#,##0'
    if widths:
        for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
    return r + 2

def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = T
    if sub: ws['A2'] = sub; ws['A2'].font = N
